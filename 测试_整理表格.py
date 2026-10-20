# -*- coding: utf-8 -*-
"""整理表格的测试。"""
import os
import tempfile
import unittest

import openpyxl

from 整理表格 import 读取名称列, 写出结果
from 归一引擎 import 归一分档


def _造表格(路径, 表名, 行们):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 表名
    for 行 in 行们:
        ws.append(行)
    wb.save(路径)


class 测试读取名称列(unittest.TestCase):
    def setUp(self):
        self.目录 = tempfile.mkdtemp()
        self.路径 = os.path.join(self.目录, "测试表.xlsx")
        _造表格(self.路径, "Sheet1",
                [["凭证号", "供应商"], ["记-001", "上海语荻光电有限公司"], ["记-002", "柏油"], ["记-003", ""]])

    def test_正常读取跳过空值(self):
        结果 = 读取名称列(self.路径, "供应商")
        self.assertEqual(结果, [(2, "上海语荻光电有限公司"), (3, "柏油")])

    def test_列名不存在报中文错(self):
        with self.assertRaises(ValueError) as 上下文:
            读取名称列(self.路径, "客户")
        self.assertIn("没有名为", str(上下文.exception))
        self.assertIn("供应商", str(上下文.exception))

    def test_文件不存在报中文错(self):
        with self.assertRaises(FileNotFoundError) as 上下文:
            读取名称列(os.path.join(self.目录, "不存在.xlsx"), "供应商")
        self.assertIn("找不到文件", str(上下文.exception))


class 测试写出结果(unittest.TestCase):
    def setUp(self):
        self.目录 = tempfile.mkdtemp()
        self.路径 = os.path.join(self.目录, "账表.xlsx")
        _造表格(self.路径, "Sheet1",
                [["凭证号", "供应商", "摘要"],
                 ["记-001", "上海语荻　光电有限公司", "销售货物"],
                 ["记-002", "上海语荻光电有限公司", "销售货物"],
                 ["记-003", "成都博立叶光学仪器有限公司", "采购26512000000000000001"],
                 ["记-004", "成都傅立叶光学仪器有限公司", "采购26512000000000000001"],
                 ["记-005", "天津满运软件科技有限公司", "运费"],
                 ["记-006", "宿迁满运软件科技有限公司", "运费"],
                 ["记-007", "入库（河北雄安星煜电子科技有限公司", "材料"],
                 ["记-008", "YF0001_成都精睿恒光电科技有限公司", "货款"],
                 ["记-009", "柏油", "入库"]])
        self.名称们 = 读取名称列(self.路径, "供应商")
        self.分档 = 归一分档([名 for _, 名 in self.名称们])
        with open(self.路径, "rb") as f:
            self.原字节 = f.read()

    def _全表文本(self, 路径):
        wb = openpyxl.load_workbook(路径)
        return "".join(str(c.value) for ws in wb.worksheets for 行 in ws.iter_rows() for c in 行 if c.value)

    def test_写出结果基本行为(self):
        出 = os.path.join(self.目录, "账表_归一后.xlsx")
        实际 = 写出结果(self.路径, "供应商", self.名称们, self.分档,
                     {1: "合", 2: "不合"}, 出, 摘要列="摘要")
        self.assertEqual(实际, 出)
        self.assertTrue(os.path.exists(出))
        with open(self.路径, "rb") as f:
            self.assertEqual(f.read(), self.原字节)  # 原文件一字未动
        wb = openpyxl.load_workbook(出)
        表 = wb["Sheet1"]
        self.assertEqual(表.cell(1, 4).value, "规范名称")
        self.assertEqual(表.cell(2, 4).value, "上海语荻光电有限公司")        # 自动归一
        self.assertEqual(表.cell(4, 4).value, "成都傅立叶光学仪器有限公司")  # 确认合，同长取字典序靠前
        self.assertEqual(表.cell(6, 4).value, "天津满运软件科技有限公司")    # 确认不合，保持原名
        文本 = self._全表文本(出)
        self.assertIn("规范名称", 文本)
        self.assertIn("原始称谓", 文本)
        self.assertIn("涉及行数", 文本)
        self.assertIn("归一依据", 文本)
        self.assertIn("同发票号", 文本)   # 博/傅两行摘要里的同一发票号
        self.assertNotIn("共现", 文本)    # 同凭证共现已彻底删除
        self.assertIn("天津满运软件科技有限公司", 文本)  # 不合的进存疑清单

    def test_候选超50对先请示不落盘(self):
        名单 = [f"统一样本{i}公司" for i in range(12)]
        分档 = 归一分档(名单)
        self.assertGreater(len(分档["候选"]), 50)
        出 = os.path.join(self.目录, "超量.xlsx")
        with self.assertRaises(ValueError) as 上下文:
            写出结果(self.路径, "供应商", [(i + 2, n) for i, n in enumerate(名单)], 分档, {}, 出)
        self.assertIn("50", str(上下文.exception))
        self.assertFalse(os.path.exists(出))

    def test_同输入同决定结果一致(self):
        出1 = os.path.join(self.目录, "一遍.xlsx")
        出2 = os.path.join(self.目录, "两遍.xlsx")
        写出结果(self.路径, "供应商", self.名称们, self.分档, {1: "合", 2: "不合"}, 出1, 摘要列="摘要")
        写出结果(self.路径, "供应商", self.名称们, self.分档, {1: "合", 2: "不合"}, 出2, 摘要列="摘要")
        self.assertEqual(self._全表文本(出1), self._全表文本(出2))

    def test_单名清洗也写入规范名称(self):
        # “入库（河北…”这类单个名字清洗后应直接作为规范名称，而不是原样返回脏名
        出 = os.path.join(self.目录, "单名清洗.xlsx")
        写出结果(self.路径, "供应商", self.名称们, self.分档, {1: "合", 2: "不合"}, 出, 摘要列="摘要")
        wb = openpyxl.load_workbook(出)
        表 = wb["Sheet1"]
        self.assertEqual(表.cell(8, 4).value, "河北雄安星煜电子科技有限公司")
        self.assertEqual(表.cell(9, 4).value, "成都精睿恒光电科技有限公司")


if __name__ == "__main__":
    unittest.main()
