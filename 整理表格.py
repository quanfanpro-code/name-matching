# -*- coding: utf-8 -*-
"""整理表格：读写表格文件，调用归一引擎，永不覆盖原文件。"""
import csv
import os
import re
from collections import Counter

import openpyxl

from 归一引擎 import 清洗名称


def 读取名称列(文件路径, 列名, 表名=None):
    """读出指定列的（行号, 名称），行号从 1 开始含表头计数，空值跳过。"""
    if not os.path.exists(文件路径):
        raise FileNotFoundError("找不到文件：" + str(文件路径))
    if str(文件路径).lower().endswith(".csv"):
        return _读csv(文件路径, 列名)
    return _读xlsx(文件路径, 列名, 表名)


def _读xlsx(文件路径, 列名, 表名):
    wb = openpyxl.load_workbook(文件路径, read_only=True, data_only=True)
    ws = wb[表名] if 表名 else wb.worksheets[0]
    行们 = ws.iter_rows(values_only=True)
    表头 = next(行们, None)
    if not 表头:
        wb.close()
        raise ValueError("表格是空的，读不到表头：" + str(文件路径))
    列名列表 = [str(x).strip() if x is not None else "" for x in 表头]
    if 列名 not in 列名列表:
        wb.close()
        raise ValueError(f"表格里没有名为「{列名}」的列，实际有的列是：" + "、".join(x for x in 列名列表 if x))
    列号 = 列名列表.index(列名)
    结果 = []
    for 行号, 行 in enumerate(行们, start=2):
        if 列号 < len(行) and 行[列号] is not None and str(行[列号]).strip():
            结果.append((行号, str(行[列号]).strip()))
    wb.close()
    return 结果


def _读csv(文件路径, 列名):
    for 编码 in ("utf-8-sig", "gbk"):
        try:
            with open(文件路径, newline="", encoding=编码) as f:
                行们 = list(csv.reader(f))
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("表格文件读不出来（编码不对）：" + str(文件路径))
    if not 行们:
        raise ValueError("表格是空的，读不到表头：" + str(文件路径))
    表头 = [x.strip() for x in 行们[0]]
    if 列名 not in 表头:
        raise ValueError(f"表格里没有名为「{列名}」的列，实际有的列是：" + "、".join(x for x in 表头 if x))
    列号 = 表头.index(列名)
    return [(行号, 行[列号].strip()) for 行号, 行 in enumerate(行们[1:], start=2)
            if 列号 < len(行) and 行[列号].strip()]


def _发票号(文本):
    return set(re.findall(r"\d{8,}", str(文本 or "")))


def 写出结果(原文件路径, 列名, 名称们, 分档结果, 确认决定, 输出路径, 摘要列=None):
    """写出结果表（加规范名称列）和归一对照表，存到新文件，永不覆盖原文件。"""
    if os.path.abspath(输出路径) == os.path.abspath(原文件路径):
        raise ValueError("输出路径不能与原文件相同（永不覆盖原文件）")
    候选 = 分档结果["候选"]
    if len(候选) > 50:
        raise ValueError(f"候选有 {len(候选)} 对，超过 50 对，请先确认数据质量后再运行")
    映射 = {}
    for 规范, 原名们 in 分档结果["自动归一"].items():
        for 原 in 原名们:
            映射[原] = 规范
    # 读原文摘要，供账内证据（同发票号）使用
    摘要按名 = {}
    wb = openpyxl.load_workbook(原文件路径)
    ws = wb.worksheets[0]
    表头 = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    摘要列号 = 表头.index(摘要列) + 1 if (摘要列 and 摘要列 in 表头) else None
    for 行号, 名 in 名称们:
        if 摘要列号:
            摘要按名.setdefault(名, _发票号(ws.cell(行号, 摘要列号).value))
        else:
            摘要按名.setdefault(名, set())
    对照 = []
    for 编号, (甲, 乙, 说明) in enumerate(候选, start=1):
        决 = 确认决定.get(编号, 确认决定.get(str(编号), ""))
        if 决 == "合":
            规范 = sorted([清洗名称(甲), 清洗名称(乙)], key=lambda s: (-len(s), s))[0]
            映射[甲] = 映射[乙] = 规范
            共同票 = 摘要按名.get(甲, set()) & 摘要按名.get(乙, set())
            依据 = "确认合并；" + (f"同发票号：{sorted(共同票)[0]}" if 共同票 else "无同发票号佐证")
            对照.append((规范, f"{甲}；{乙}", 说明, 依据))
        else:
            对照.append((f"【存疑{编号}】{甲} ／ {乙}", "", 说明,
                         "确认不合" if 决 == "不合" else "未确认，保持原名"))
    # 汇总涉及行数
    计数 = Counter(名 for _, 名 in 名称们)
    # 写结果表：原表加“规范名称”列
    列号 = 表头.index(列名) + 1
    新列 = ws.max_column + 1
    ws.cell(1, 新列, "规范名称")
    for 行号, 名 in 名称们:
        ws.cell(行号, 新列, 映射.get(名, 清洗名称(名)))
    # 写归一对照表
    对照表 = wb.create_sheet("归一对照")
    对照表.append(["规范名称", "原始称谓", "涉及行数", "归一依据"])
    for 规范, 原名们 in sorted(分档结果["自动归一"].items()):
        for 原 in 原名们:
            对照表.append([规范, 原, 计数[原], "清洗后一致或简称包含，机器自动归一"])
    for 规范, 原述, 说明, 依据 in 对照:
        对照表.append([规范, 原述, "", f"{说明}；{依据}"])
    wb.save(输出路径)
    wb.close()
    return 输出路径
